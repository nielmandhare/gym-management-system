"""Read the supplied Word report without changing it."""
from pathlib import Path
from docx import Document

root = Path(__file__).resolve().parents[2]
doc = Document(root / 'Gym_Management_System_STQA_Report_Selenium_IDE.docx')
out = root / 'report-work'
out.mkdir(parents=True, exist_ok=True)
lines = []
for i, p in enumerate(doc.paragraphs):
    lines.append(f'P{i} [{getattr(p.style, "name", "None")}] {p.text}')
for i, table in enumerate(doc.tables):
    lines.append(f'\nTABLE {i}')
    for row in table.rows:
        lines.append(' | '.join(c.text.replace('\n', ' / ') for c in row.cells))
(out / 'original-text.txt').write_text('\n'.join(lines), encoding='utf-8')
print('Sections', len(doc.sections))
for s in doc.sections:
    print('Page inches', s.page_width.inches, s.page_height.inches,
          'Margins', s.top_margin.inches, s.bottom_margin.inches,
          s.left_margin.inches, s.right_margin.inches)
print('Styles', [(s.name, s.font.name, s.font.size.pt if s.font.size else None)
                 for s in doc.styles if s.name in ('Normal', 'Heading 1', 'Heading 2', 'Title')])
print('Text extraction saved')
