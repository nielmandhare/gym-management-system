"""Build the gym SRS using the supplied friend's PDF as a visual reference."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'STQA_Submission' / '01_Report' / 'Gym_Management_System_SRS_Final.docx'


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.8)
    sec.footer_distance = Inches(.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12)
        fonts = style.element.get_or_add_rPr().rFonts
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():
                del fonts.attrib[attr]
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attr), 'Times New Roman')
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
    doc.styles['Title'].font.size = Pt(18)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Title'].paragraph_format.space_after = Pt(16)
    doc.styles['Subtitle'].font.size = Pt(13)
    doc.styles['Subtitle'].font.bold = True
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Subtitle'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Subtitle'].paragraph_format.space_after = Pt(12)
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].font.bold = True
        doc.styles[name].font.size = Pt(16 if name == 'Heading 1' else 14)
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.space_after = Pt(8)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles['List Bullet'].paragraph_format.space_after = Pt(3)
    doc.styles['List Bullet'].paragraph_format.left_indent = Inches(.25)
    doc.styles['List Bullet'].paragraph_format.first_line_indent = Inches(-.13)
    doc.core_properties.title = 'Gym Management System Software Requirements Specification'
    doc.core_properties.subject = 'STQA Assignment 5'
    doc.core_properties.author = ''

    def p(text):
        para = doc.add_paragraph(text)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return para

    def h(text, level=2):
        return doc.add_heading(text, level)

    def bullet(label, text=''):
        para = doc.add_paragraph(style='List Bullet')
        if text:
            para.add_run(label + ': ').bold = True
            para.add_run(text)
        else:
            para.add_run(label)
        return para

    # Page 1 follows the reference's title, problem statement and opening sections.
    doc.add_paragraph('Software Requirements Specification', 'Title')
    doc.add_paragraph('Gym Management System', 'Subtitle')
    h('Problem statement', 2)
    p('To provide a web-based Gym Management System that allows administrators, trainers and members to access related gym records through one application. The system shall simplify member administration, attendance and fitness tracking, maintain consistent records, and restrict each user to the functions and information permitted for that role.')
    h('1. Introduction', 1)
    h('1.1 Purpose')
    p('The purpose of the system is to support daily gym administration and member self-service through a browser. Admin manages gym records and accounts, Trainer reviews assigned members and enters fitness progress, and Member views personal records. This SRS defines the functions, operating environment and interfaces used as a basis for development and STQA testing.')
    h('1.2 Product Scope')
    p('The system shall provide the following major functions:')
    for text in [
        'Login, logout and role-based access for Admin, Trainer and Member.',
        'Member and trainer profiles and user account management.',
        'Membership details, paid/due status and member attendance.',
        'Assigned-member views, measurements, BMI and progress history.',
        'Equipment and product records, including condition, stock and price.',
        'Workout plan details and existing payment, sales and training history.',
        'Search, filtered CSV exports, six PDF reports and local SQL backup.',
    ]:
        bullet(text)
    p('The current web scope includes viewing payment and sales transactions and training assignments. New transaction entry and creation of training assignments are outside this version.')

    doc.add_page_break()
    h('1.3 Overview')
    p('The application consists of a browser interface, a Java server and a MySQL database. Users access a common login page and then receive role-specific navigation. This document is intended for the project team, faculty, gym users and testers. Section 1 introduces the project, Section 2 describes its functions and constraints, and Section 3 defines its external interfaces.')
    h('1.4 References')
    bullet('STQA Lab Manual, BE COMP, Sem I, MMCOE Pune, Assignment 5, PDF pages 26-29; supplied project copy in output/stqa.')
    bullet('Gym Management System web edition documentation, web/README.md, current project checkout.')
    bullet('Gym Management System OOMD Mini Project Report, Gym_Management_System_OOMD_Report.pdf, project root.')
    h('2. Overall Description', 1)
    h('2.1 Product Perspective')
    p('The web application is a browser alternative to the existing Java Swing gym project. It uses the existing database schema and three user roles. Its major components are:')
    bullet('Frontend', 'HTML, CSS and JavaScript screens for users and forms.')
    bullet('Backend', 'Java HTTP server for requests, validation and role checks.')
    bullet('Database', 'MySQL stores accounts and related gym records through JDBC.')
    h('2.2 Product Functions')
    bullet('User management', 'Admin provisions accounts, assigns roles and controls active status; users may change their own passwords.')
    bullet('Gym administration', 'Admin maintains profiles, memberships, attendance, equipment and products, with guarded deletion of linked records.')
    bullet('Trainer functions', 'View assigned members and record validated measurements for active assignments; calculate BMI from height and weight.')
    bullet('Member functions', 'View personal membership, attendance, measurements and training history; update permitted profile details.')
    bullet('Records and reports', 'Search stored records, view workout details, export CSV, download operational PDFs and create an Admin-only SQL backup.')

    doc.add_page_break()
    h('2.3 User Classes and Characteristics')
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.1), Inches(5.8)]
    table.columns[0].width, table.columns[1].width = widths
    table.rows[0].cells[0].text = 'User Class'
    table.rows[0].cells[1].text = 'Characteristics'
    for label, text in [
        ('Admin', 'Maintains gym records, users and permissions and accesses operational reports and backup.'),
        ('Trainer', 'Regularly reviews assigned members and records their fitness measurements and progress.'),
        ('Member', 'Uses self-service screens to view personal membership, attendance, progress and training history.'),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, text
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + name)
        for k, v in {'val': 'single', 'sz': '4', 'color': '000000'}.items():
            edge.set(qn('w:' + k), v)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for ri, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, cell in enumerate(row.cells):
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                element = OxmlElement('w:' + side)
                element.set(qn('w:w'), '80')
                element.set(qn('w:type'), 'dxa')
                margins.append(element)
            cell._tc.get_or_add_tcPr().append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    run.bold = ri == 0 or ci == 0
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p('All users are expected to have basic browser and form-entry skills.')
    h('2.4 Operating Environment')
    bullet('Operating system', 'Windows with PowerShell for the supplied setup scripts.')
    bullet('Client', 'Chrome or Edge browser.')
    bullet('Frontend', 'HTML, CSS and JavaScript.')
    bullet('Backend', 'JDK 17 or newer; Java built-in HTTP server and JDBC.')
    bullet('Database', 'MySQL Server 8.x on the dedicated local port 3307.')
    bullet('Application', 'Local browser access at http://127.0.0.1:8081.')
    bullet('Testing', 'Selenium WebDriver in Java/Python and Selenium IDE.')
    h('2.5 Design and Implementation Constraints')
    for text in [
        'The server shall enforce active accounts, linked profiles and role permissions; passwords shall be stored as hashes.',
        'The browser shall access the database through the backend. Database operations shall use JDBC and parameterised SQL.',
        'Forms shall validate inputs, prevent duplicate attendance and maintain related data through transactions and guarded deletion.',
        'The interface shall support readable forms, clear errors and narrower browser layouts.',
        'The current setup is a local Windows demonstration. Public hosting, online payments and biometric integration are outside its scope.',
    ]:
        bullet(text)

    doc.add_page_break()
    h('3. External Interface Requirements', 1)
    h('3.1 User Interfaces')
    p('The application shall provide the following interfaces:')
    bullet('Login', 'Common sign-in and clear validation messages.')
    bullet('Admin dashboard', 'Gym totals, management tables, record forms, reports and user controls.')
    bullet('Trainer dashboard', 'Own profile, assigned members, attendance views and progress entry.')
    bullet('Member dashboard', 'Personal profile, membership, attendance, progress and training history.')
    bullet('Record views', 'Search, pagination, applicable CSV exports and workout details.')
    p('Navigation and controls shall be consistent. Users shall receive validation feedback, appropriate confirmations and access to logout.')
    h('3.2 Hardware Interfaces')
    p('The system shall use a standard computer, display, keyboard and mouse or touch input. It requires local storage for database files and backups. A dual-core computer with 8 GB RAM is a suggested classroom setup, not a measured minimum. No specialised attendance reader or fitness-machine connection is required.')
    h('3.3 Software Interfaces')
    bullet('Browser and Java server', 'Forms send requests and receive JSON responses or downloadable files.')
    bullet('MySQL', 'Connector/J 8.4.0 provides JDBC access to account and gym data in MySQL 8.x.')
    bullet('Reports and backup', 'OpenPDF 1.3.30 generates the members, trainers, attendance, payments, sales and equipment PDFs. The backup utility writes a local SQL file.')
    bullet('Testing tools', 'Selenium WebDriver and Selenium IDE operate the browser to verify application behaviour.')
    h('3.4 Communications Interfaces')
    bullet('HTTP', 'The browser communicates with the application over the loopback connection on port 8081.')
    bullet('TCP/IP', 'The backend connects to the dedicated MySQL instance at 127.0.0.1:3307.')
    bullet('Data and sessions', 'The application uses form submissions, JSON responses and a session cookie. The server checks access and validates state-changing requests.')
    p('The local demonstration uses HTTP without TLS. Email, SMS and external service integrations are not required.')

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run('Gym Management System | Software Requirements Specification | Page ')
    run.font.name = 'Arial'
    run.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
